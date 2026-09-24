"""core/procesador.py — Lógica de extracción de PDFs y generación de Excel"""
import re
import csv
import shutil
import openpyxl
import pdfplumber
import io
from pathlib import Path
from typing import List, Dict, Optional, Tuple, Union
from .config import (
    logger, SOL_DIR, RUTA_PLANTILLA, RUTA_PLANTILLA_SALIDA,
    RUTA_PLANTILLA_ENTRADA, RUTA_ACTIVOS, MAPEO_ACTIVOS
)
from .catalogo import obtener_nombre_compania, leer_activos
from .db import guardar_solicitud
from datetime import datetime


# ──────────────────────────────────────────────────────────────
#  Utilidades
# ──────────────────────────────────────────────────────────────

def es_int(valor: str) -> bool:
    try:
        int(valor)
        return True
    except (ValueError, TypeError):
        return False


# ──────────────────────────────────────────────────────────────
#  Excel helpers
# ──────────────────────────────────────────────────────────────

def quitar_celdas_combinadas(hoja) -> None:
    if hoja.merged_cells:
        for rango in list(hoja.merged_cells.ranges):
            hoja.unmerge_cells(str(rango))


def limpiar_celdas(hoja, celda_inicio: str) -> None:
    start_row = hoja[celda_inicio].row
    start_col = hoja[celda_inicio].column
    for row in hoja.iter_rows(
        min_row=start_row, min_col=start_col,
        max_row=hoja.max_row, max_col=hoja.max_column
    ):
        for cell in row:
            cell.value = None


# ──────────────────────────────────────────────────────────────
#  Extracción de PDF
# ──────────────────────────────────────────────────────────────

def extraer_personal_sube(ruta: Path) -> List[Dict]:
    """Extrae filas 1-S del PDF."""
    personal = []
    try:
        with pdfplumber.open(ruta) as pdf:
            numero = 1
            for pagina in pdf.pages:
                for registro in pagina.extract_tables():
                    for celda_original in registro:
                        celda = [x for x in celda_original if x is not None]
                        if re.search(r"- S", str(celda)):
                            personal.append({
                                "id":       numero,
                                "rfc":      celda[1].lstrip("0") if len(celda) > 1 else "",
                                "nombre":   celda[2]             if len(celda) > 2 else "",
                                "libreta":  celda[3][:10]        if len(celda) > 3 else "",
                                "vigencia": celda[4][:10]        if len(celda) > 4 else "",
                                "llegada":  celda[5]             if len(celda) > 5 else "",
                                "salida":   celda[6]             if len(celda) > 6 else "",
                                "dias":     celda[7]             if len(celda) > 7 else "",
                            })
                            numero += 1
        logger.info(f"Extraídos {len(personal)} registros (sube) de {ruta.name}")
    except Exception as e:
        logger.error(f"Error extrayendo personal de {ruta}: {e}")
        raise
    return personal


def extraer_personal_baja(ruta: Path) -> List[Dict]:
    """Extrae filas 1-B del PDF."""
    personal = []
    try:
        with pdfplumber.open(ruta) as pdf:
            numero = 1
            for pagina in pdf.pages:
                for registro in pagina.extract_tables():
                    for celda_original in registro:
                        celda = [x for x in celda_original if x is not None]
                        if re.search(r"- B", str(celda)):
                            personal.append({
                                "id":       numero,
                                "rfc":      celda[1].lstrip("0") if len(celda) > 1 else "",
                                "nombre":   celda[2]             if len(celda) > 2 else "",
                                "libreta":  celda[3][:10]        if len(celda) > 3 else "",
                                "vigencia": celda[4][:10]        if len(celda) > 4 else "",
                                "llegada":  celda[5]             if len(celda) > 5 else "",
                                "salida":   celda[6]             if len(celda) > 6 else "",
                                "dias":     celda[7]             if len(celda) > 7 else "",
                            })
                            numero += 1
    except Exception as e:
        logger.error(f"Error extrayendo bajas de {ruta}: {e}")
    return personal


def extraer_datos_solicitud(ruta: Path) -> Dict:
    """Extrae metadatos del encabezado del PDF."""
    info: Dict = {"nombre_compania": "SIN NOMBRE", "destinohosp": "None",
                  "activoA": False, "activoS": False}
    try:
        with pdfplumber.open(ruta) as pdf:
            primera = pdf.pages[0]
            texto   = primera.extract_text()
            lineas  = texto.splitlines()
            tablas  = primera.extract_tables()

            # Número de solicitud
            if lineas:
                partes = lineas[0].split()
                info["solicitud"] = partes[2] if len(partes) > 2 else "SIN_NUM"

            # Contrato (líneas 8-9)
            for li in [8, 9]:
                if li < len(lineas):
                    for palabra in lineas[li].split():
                        if es_int(palabra) and len(palabra) > 8:
                            info["contrato"] = palabra
                            break

            # Destino
            info["destinohosp"] = "RPX" if "HOSP: RPX" in texto else "None"

            # Activos
            cadena = " ".join(lineas[5:7]) if len(lineas) > 6 else ""
            lista_activos = leer_activos()
            a1, a2 = _buscar_activos(cadena, lista_activos)
            info["activoA"] = a1 or False
            info["activoS"] = a2 or False

            # Códigos presupuestales
            for codigo, prefijo in [
                ("elepep", "ELE:"), ("progpre", "PRO:"),
                ("cge", "CGE:"), ("cta", "CTA:"), ("pos", "POS:")
            ]:
                idx = texto.find(prefijo)
                if idx != -1:
                    lng = 16 if codigo in ["elepep", "progpre"] else 8 if codigo != "pos" else 9
                    info[codigo] = texto[idx + 4: idx + 4 + lng]

            # Firmas
            if len(tablas) > 3 and tablas[3]:
                fila0 = tablas[3][0]
                if fila0:
                    ls = (fila0[0] or "").splitlines()
                    info["nombre_solicita"] = ls[1] if len(ls) > 1 else ""
                    info["ficha_solicita"]  = ls[2] if len(ls) > 2 else ""
                if len(fila0) > 1:
                    la = (fila0[1] or "").splitlines()
                    info["nombre_autoriza"] = la[1] if len(la) > 1 else ""
                    info["ficha_autoriza"]  = la[2] if len(la) > 2 else ""

            # Tipo / razón social / transporte
            clave = "Tipo de Personal: Compañía: No. Contrato SAP: Transporte: Itinerario ida: Fec/Hora Itinerario:"
            try:
                idx = lineas.index(clave)
                if idx + 1 < len(lineas):
                    linea_datos = lineas[idx + 1]
                    tokens = linea_datos.split()
                    info["tipo"] = tokens[0] if tokens else ""

                    if info["tipo"] == "PEMEX":
                        info["contrato"]    = "N/A"
                        info["razonsocial"] = "PEMEX"
                    else:
                        contrato_enc = ""
                        for t in tokens[1:]:
                            try:
                                int(t)
                                if len(t) > 8:
                                    contrato_enc = t
                                    break
                            except ValueError:
                                pass
                        if contrato_enc:
                            info["contrato"]    = contrato_enc
                            i = linea_datos.find(contrato_enc)
                            info["razonsocial"] = linea_datos[len(info["tipo"]):i].strip()
                        elif "contrato" in info:
                            i = linea_datos.find(info["contrato"])
                            if i != -1:
                                info["razonsocial"] = linea_datos[len(info["tipo"]):i].strip()
            except ValueError:
                pass

            # Transporte
            for t in ["MEDIOS PROPIOS", "MARITIMO", "AEREO"]:
                for ln in lineas:
                    if t in ln.upper():
                        info["transporte"] = t
                        break
                if "transporte" in info:
                    break

            logger.info(f"Datos extraídos de solicitud {info.get('solicitud','?')}")
    except Exception as e:
        logger.error(f"Error extrayendo datos de {ruta}: {e}")
        raise
    return info


def _buscar_activos(texto: str, lista: List[str]) -> Tuple[Optional[str], Optional[str]]:
    encontrados = []
    for activo in lista:
        inicio = 0
        while True:
            i = texto.find(activo, inicio)
            if i == -1:
                break
            encontrados.append((i, activo))
            inicio = i + len(activo)
    encontrados.sort(key=lambda x: x[0])
    def expandir(a):
        return MAPEO_ACTIVOS.get(a, a) if a else None
    a1 = expandir(encontrados[0][1]) if len(encontrados) >= 1 else None
    a2 = expandir(encontrados[1][1]) if len(encontrados) >= 2 else None
    return a1, a2


# ──────────────────────────────────────────────────────────────
#  Generador de plantilla principal (solicitud → xlsx)
# ──────────────────────────────────────────────────────────────

def llenar_plantilla(ruta_pdf: Path) -> Tuple[bool, str]:
    """Procesa un PDF y genera el Excel de solicitud."""
    try:
        if not RUTA_PLANTILLA.exists():
            return False, f"Plantilla no encontrada: {RUTA_PLANTILLA}"
        if not ruta_pdf.exists():
            return False, f"PDF no encontrado: {ruta_pdf}"

        personal = extraer_personal_sube(ruta_pdf)
        info     = extraer_datos_solicitud(ruta_pdf)

        if not personal:
            return False, "No se encontró personal en el PDF"

        personal_baja = extraer_personal_baja(ruta_pdf)

        wb   = openpyxl.load_workbook(RUTA_PLANTILLA)
        hoja = wb.active

        quitar_celdas_combinadas(hoja)
        limpiar_celdas(hoja, "A15")

        fila_inicial = 15
        for i, p in enumerate(personal, start=fila_inicial):
            hoja[f"A{i}"] = p["id"]
            hoja[f"C{i}"] = p["rfc"]
            hoja[f"D{i}"] = p["nombre"]
            hoja[f"G{i}"] = p["libreta"]
            hoja[f"H{i}"] = p["vigencia"]
            hoja[f"I{i}"] = p["llegada"]
            hoja[f"J{i}"] = p["salida"]
            hoja[f"K{i}"] = p["dias"]

        fila_firma = len(personal) + fila_inicial + 3
        hoja[f"C{fila_firma}"].value   = info.get("nombre_solicita", "")
        hoja[f"C{fila_firma+1}"].value = info.get("ficha_solicita", "")
        hoja[f"I{fila_firma}"].value   = info.get("nombre_autoriza", "")
        hoja[f"I{fila_firma+1}"].value = info.get("ficha_autoriza", "")

        hoja["A11"].value = info.get("progpre", "")
        hoja["J1"].value  = info.get("solicitud", "")
        hoja["F11"].value = info.get("elepep", "")
        hoja["G11"].value = "PEF"
        hoja["H11"].value = info.get("pos", "")
        hoja["I11"].value = info.get("cge", "")
        hoja["J11"].value = info.get("cta", "")
        hoja["A8"].value  = info.get("tipo", "")
        hoja["G8"].value  = info.get("contrato", "")

        if "razonsocial" in info:
            hoja["E8"].value = obtener_nombre_compania(info["razonsocial"])

        errores = []
        if info["activoA"]:
            hoja["E6"].value = info["activoA"]
        else:
            errores.append("Activo solicitante no encontrado")
        if info["activoS"]:
            hoja["H6"].value = info["activoS"]
        else:
            errores.append("Activo autorizador no encontrado")

        nombre_salida = f"{info.get('solicitud', 'SIN_NUM')}.xlsx"
        ruta_salida   = SOL_DIR / nombre_salida
        wb.save(ruta_salida)

        # Guardar metadatos en SQLite para listado rápido
        compania_corta = obtener_nombre_compania(info.get("razonsocial", "")) \
                         if "razonsocial" in info else ""
        info["compania_corta"] = compania_corta
        guardar_solicitud(info, personal, personal_baja, nombre_salida)

        mensaje = f"Archivo generado: {nombre_salida}"
        if personal_baja:
            mensaje += f" | Bajas: {len(personal_baja)} persona(s)"
        if errores:
            mensaje += f" | Advertencias: {', '.join(errores)}"

        logger.info(mensaje)
        return True, mensaje

    except Exception as e:
        msg = f"Error procesando {ruta_pdf.name}: {e}"
        logger.error(msg)
        return False, msg


def listar_solicitudes_xlsx(page: int = 1, limit: int = 20, search: Optional[str] = None) -> tuple[list[dict], int]:
    """
    Lee solicitudes desde Supabase con paginación y búsqueda opcional.
    Devuelve una tupla: (lista de dicts, total_de_registros).
    """
    from .supabase_db import _get
    
    try:
        offset = (page - 1) * limit
        
        # Construir filtros
        filters = {}
        if search:
            # Supabase usa 'ilike' para búsqueda parcial insensible a mayúsculas
            filters["numero"] = f"ilike.%{search.strip()}%"
        
        # 1. Obtener los registros paginados
        params = {
            "select": "*", 
            "order": "fecha_llegada.desc,procesado.desc", 
            "limit": limit, 
            "offset": offset
        }
        
        # Combinar filtros y parámetros
        query_params = {**params}
        if filters:
            query_params.update(filters)
            
        rows = _get("solicitudes", filters=query_params)
        
        # 2. Obtener el total de registros para calcular las páginas
        # Si hay búsqueda, el total debe reflejar solo los resultados filtrados
        total_params = {"select": "id"}
        if search:
            total_params["numero"] = f"ilike.%{search.strip()}%"
            
        total_rows = _get("solicitudes", **total_params) 
        total_count = len(total_rows)

        result = []
        for r in rows:
            fecha_mod = "" 
            
            result.append({
                "id":          r["id"],
                "nombre":      r["numero"],
                "compania":    r["compania"] or "",
                "n":           r["n_personas"],
                "fecha":       r["fecha_llegada"] or "",
                "fecha_mod":   fecha_mod,
                "procesado":   r.get("procesado") or "",
                "tipo":        r.get("tipo") or "",
                "archivo":     r.get("archivo") or "",
                "tiene_bajas": bool(r["tiene_bajas"]),
                "n_bajas":     r["n_bajas"],
                "destino":     r.get("destinohosp") or "",
            })
        return result, total_count
    except Exception as e:
        logger.error(f"Error en listar_solicitudes_xlsx: {e}")
        return [], 0


def leer_meta_xlsx(ruta: Path) -> dict:
    """Fallback: abre el xlsx directamente. Solo se usa para archivos
    que aún no tienen registro en la BD (solicitudes antiguas)."""
    meta = {"compania": "", "n": 0, "fecha": ""}
    try:
        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        ws = wb.active
        meta["compania"] = str(ws["E8"].value or "").strip()
        n = 0
        fecha = ""
        for row in ws.iter_rows(min_row=15, max_row=ws.max_row,
                                min_col=4, max_col=9, values_only=True):
            if not row[0] or not str(row[0]).strip():
                break
            n += 1
            if not fecha and row[5]:
                fecha = str(row[5]).split(" ")[0][:10]
        meta["n"]     = n
        meta["fecha"] = fecha
        wb.close()
    except Exception:
        pass
    return meta


# ──────────────────────────────────────────────────────────────
#  Plantilla de salidas (bajas)
# ──────────────────────────────────────────────────────────────

_FILA_INICIO_SAL  = 22
_POR_BLOQUE_SAL   = 15
_PASO_SAL         = 17

def _fila_salida(idx: int) -> int:
    b = idx // _POR_BLOQUE_SAL
    p = idx %  _POR_BLOQUE_SAL
    return _FILA_INICIO_SAL + (0 if b == 0 else 17 + (b-1)*_PASO_SAL) + p


def llenar_plantilla_salidas(registros: List[Dict], col_nombre: str,
                              col_depto: Optional[str], col_cama: Optional[str],
                              col_solicitud: Optional[str], col_transporte: Optional[str],
                              ruta_destino: Path) -> int:
    if not RUTA_PLANTILLA_SALIDA.exists():
        raise FileNotFoundError(f"Plantilla no encontrada: {RUTA_PLANTILLA_SALIDA}")
    registros = sorted(registros, key=lambda f: (
        (f.get(col_depto)     or "").upper(),
        (f.get(col_solicitud) or "").upper(),
        (f.get(col_nombre)    or "").upper(),
    ))
    shutil.copy2(RUTA_PLANTILLA_SALIDA, ruta_destino)
    wb = openpyxl.load_workbook(ruta_destino)
    ws = wb["SALIDAS"]
    for idx, fila in enumerate(registros):
        r = _fila_salida(idx)
        ws.cell(r, 3).value  = (fila.get(col_nombre)     or "").strip()
        ws.cell(r, 4).value  = (fila.get(col_depto)      or "").strip() if col_depto     else ""
        ws.cell(r, 6).value  = (fila.get(col_solicitud)  or "").strip() if col_solicitud else ""
        ws.cell(r, 7).value  = (fila.get(col_cama)       or "").strip() if col_cama      else ""
        ws.cell(r, 8).value  = (fila.get(col_transporte) or "").strip() if col_transporte else ""
    wb.save(ruta_destino)
    return len(registros)


# ──────────────────────────────────────────────────────────────
#  Plantilla de entradas (altas)
# ──────────────────────────────────────────────────────────────

_FILA_INICIO_ENT = 22
_POR_BLOQUE_ENT  = 15
_PASO_ENT        = 17

def _fila_entrada(idx: int) -> int:
    b = idx // _POR_BLOQUE_ENT
    p = idx %  _POR_BLOQUE_ENT
    return _FILA_INICIO_ENT + (0 if b == 0 else 17 + (b-1)*_PASO_ENT) + p


def leer_personal_de_xlsx(source: Union[Path, io.BytesIO]) -> List[Dict]:
    personal = []
    try:
        # Carga el libro desde el archivo o el stream de bytes
        wb = openpyxl.load_workbook(source, data_only=True)
        ws = wb.active
        compania   = str(ws["E8"].value or "").strip()
        transporte = ""
        for col in range(7, 12):
            v = str(ws.cell(8, col).value or "")
            for t in ["MEDIOS PROPIOS", "MARITIMO", "AEREO"]:
                if t in v.upper():
                    transporte = t
                    break
            if transporte:
                break
        
        # El numero de solicitud está en la celda J1
        num_sol = str(ws["J1"].value or "").strip()
        if not num_sol:
            # Si es un archivo Path, usamos el nombre sin extensión. Si es BytesIO, usamos un valor genérico o el esperado.
            num_sol = source.stem if isinstance(source, Path) else "MEM_FILE"

        for row in ws.iter_rows(min_row=15, max_row=ws.max_row):
            id_val = row[0].value
            if id_val is None:
                break
            try:
                int(str(id_val).strip())
            except ValueError:
                break
            rfc    = str(row[2].value or "").strip()
            nombre = str(row[3].value or "").strip()
            if not nombre and not rfc:
                break
            personal.append({"nombre": nombre, "rfc": rfc,
                              "compania": compania, "solicitud": num_sol,
                              "transporte": transporte})
    except Exception as e:
        name = source.name if isinstance(source, Path) else "BytesStream"
        logger.error(f"Error leyendo {name}: {e}")
    return personal


def llenar_plantilla_entrada(registros: List[Dict], ruta_destino: Path) -> int:
    if not RUTA_PLANTILLA_ENTRADA.exists():
        raise FileNotFoundError(f"Plantilla no encontrada: {RUTA_PLANTILLA_ENTRADA}")
    registros = sorted(registros, key=lambda r: (
        (r.get("compania")  or "").upper(),
        (r.get("solicitud") or "").upper(),
        (r.get("nombre")    or "").upper(),
    ))
    shutil.copy2(RUTA_PLANTILLA_ENTRADA, ruta_destino)
    wb = openpyxl.load_workbook(ruta_destino)
    ws = wb["ENTRADA"]
    for idx, reg in enumerate(registros):
        f = _fila_entrada(idx)
        ws.cell(f, 3).value  = reg.get("nombre",     "")
        ws.cell(f, 4).value  = ""
        ws.cell(f, 5).value  = reg.get("rfc",        "")
        ws.cell(f, 6).value  = reg.get("compania",   "")
        ws.cell(f, 7).value  = ""
        ws.cell(f, 8).value  = reg.get("solicitud",  "")
        ws.cell(f, 9).value  = ""
        ws.cell(f, 10).value = reg.get("transporte", "")
    wb.save(ruta_destino)
    return len(registros)
